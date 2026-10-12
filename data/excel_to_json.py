import json, sys
from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else "qa.xlsx"
ws = load_workbook(path, read_only=True).active

rows = list(ws.iter_rows(values_only=True))
header = [str(h).strip().lower() if h else "" for h in rows[0]]
qi, ai = header.index("instruction"), header.index("output")

data, seen, skipped = [], set(), 0
for r in rows[1:]:
    q = (r[qi] or "").strip() if isinstance(r[qi], str) else ""
    a = (r[ai] or "").strip() if isinstance(r[ai], str) else ""
    if not q or not a or (q, a) in seen:
        skipped += 1
        continue
    seen.add((q, a))
    data.append({"instruction": q, "input": "", "output": a,
                 "system": "Always reply in Dzongkha."})

json.dump(data, open("dataset.json", "w", encoding="utf-8"),
          ensure_ascii=False, indent=2)
print(f"{len(data)} pairs written, {skipped} empty or duplicate rows skipped")
